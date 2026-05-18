#!/usr/bin/env python3
"""Convert MEMBER_APIS.md to Word (.docx) for distribution."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            try:
                r.font.size = Pt(9)
            except Exception:
                pass
        else:
            paragraph.add_run(part)


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def is_table_sep(line: str) -> bool:
    s = line.strip()
    if not (s.startswith("|") and "---" in s):
        return False
    return bool(re.match(r"^\|[\s\-:|]+\|\s*$", s))


def parse_table_rows(rows: list[str]) -> tuple[list[str], list[list[str]]]:
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = []
    for row in rows[2:]:
        body.append([c.strip() for c in row.strip().strip("|").split("|")])
    return header, body


def add_table(doc: Document, rows: list[str]) -> None:
    if len(rows) < 2:
        return
    header, body = parse_table_rows(rows)
    ncols = len(header)
    nrows = 1 + len(body)
    table = doc.add_table(rows=nrows, cols=ncols)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        hdr_cells[i].text = ""
        add_inline_runs(hdr_cells[i].paragraphs[0], h)
    for ri, row in enumerate(body):
        for ci, cell in enumerate(row[:ncols]):
            table.rows[ri + 1].cells[ci].text = cell
        for ci in range(len(row), ncols):
            table.rows[ri + 1].cells[ci].text = ""


def md_to_docx(md_path: Path, out_path: Path) -> None:
    doc = Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()

    i = 0
    in_code = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        if in_code:
            if line.strip().startswith("```"):
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                try:
                    run.font.size = Pt(8)
                except Exception:
                    pass
                code_lines = []
                in_code = False
            else:
                code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("```"):
            in_code = True
            code_lines = []
            i += 1
            continue

        if is_table_row(line) and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            table_rows = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and is_table_row(lines[i]):
                table_rows.append(lines[i])
                i += 1
            add_table(doc, table_rows)
            continue

        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            level = len(m.group(1)) - 1
            title = m.group(2).strip()
            doc.add_heading(title, level=min(level, 8))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main() -> None:
    root = Path(__file__).resolve().parent.parent
    md = root / "MEMBER_APIS.md"
    out = root / "MEMBER_APIS.docx"
    if len(sys.argv) >= 2:
        md = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        out = Path(sys.argv[2])
    if not md.exists():
        print(f"Missing {md}", file=sys.stderr)
        sys.exit(1)
    md_to_docx(md, out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
